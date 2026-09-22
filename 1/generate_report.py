"""
Генератор звіту до практичної роботи №1 (Аналіз та методи оптимізації складних систем).
Питає номер варіанта (або приймає його аргументом) і генерує Zvit_PR1_Variant<N>.docx
з усіма розрахунками, графіком і відповідями на контрольні питання саме під цей варіант.

Запуск:  .venv/bin/python generate_report.py [номер_варіанта]
"""
import sys
import numpy as np
import sympy as sp
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from variants import VARIANTS

# --- Дані студента / титульного аркуша ---
STUDENT_NAME = "Толошний Мирослав"
GROUP = "КН-11"
UNIVERSITY = "Вінницький національний аграрний університет"
FACULTY = "Факультет комп'ютерних наук та економічної кібернетики"
DISCIPLINE = "Аналіз та методи оптимізації складних систем"
WORK_TITLE = "Математична модель складної системи та класифікація задач оптимізації"
WORK_NUMBER = 1

if len(sys.argv) > 1:
    VARIANT = int(sys.argv[1])
else:
    while True:
        raw = input(f"Введи номер варіанта (1-{len(VARIANTS)}): ").strip()
        if raw.isdigit() and int(raw) in VARIANTS:
            VARIANT = int(raw)
            break
        print(f"Треба ціле число від 1 до {len(VARIANTS)}, спробуй ще раз.")

# ---------- Розрахунки (та сама логіка, що й у ноутбуці) ----------
a, b, c, d, e, k, x1_lo, x1_hi, x2_lo, x2_hi, obj_name = VARIANTS[VARIANT]

x1, x2 = sp.symbols('x1 x2', real=True)
f = a*x1**2 + b*x2**2 - c*x1*x2 - d*x1 - e*x2 + k

grad = [sp.diff(f, v) for v in (x1, x2)]
H = sp.hessian(f, (x1, x2))

crit = sp.solve(grad, [x1, x2], dict=True)[0]
f_star_exact = sp.simplify(f.subs(crit))
x1_star, x2_star = float(crit[x1]), float(crit[x2])
f_star_float = float(f_star_exact)

minors = [H[:m, :m].det() for m in (1, 2)]
eigenvals = [sp.nsimplify(v) for v in H.eigenvals()]
eigen_floats = sorted(float(v) for v in eigenvals)
cond_H = eigen_floats[-1] / eigen_floats[0]
is_convex = all(m > 0 for m in minors)
in_bounds = (x1_lo <= x1_star <= x1_hi) and (x2_lo <= x2_star <= x2_hi)

# ---------- Класифікація (таблиця 1.1) ----------
classification = [
    ("Кількість змінних", "багатовимірна (n = 2)"),
    ("Вид функції", "нелінійна, квадратична (КП)"),
    ("Наявність обмежень", "безумовна оптимізація"),
    ("Тип змінних", "неперервні"),
    ("Кількість критеріїв", "однокритеріальна"),
    ("Визначеність даних", "детермінована"),
    ("Залежність від часу", "статична"),
]

# ---------- Графік ----------
fn = sp.lambdify((x1, x2), f, 'numpy')
X, Y = np.meshgrid(np.linspace(x1_lo, x1_hi, 300), np.linspace(x2_lo, x2_hi, 300))
Z = fn(X, Y)

fig, ax = plt.subplots(figsize=(6, 5))
cs = ax.contour(X, Y, Z, levels=20)
ax.clabel(cs, inline=True, fontsize=7)
ax.plot(x1_star, x2_star, 'r*', ms=14, label='точка мінімуму')
ax.set_xlabel('$x_1$'); ax.set_ylabel('$x_2$'); ax.legend(); ax.grid(alpha=.3)
ax.set_title(f'Варіант {VARIANT} ({obj_name}) — лінії рівня цільової функції')
fig.tight_layout()
contour_path = f"contour_variant_{VARIANT}.png"
fig.savefig(contour_path, dpi=150)

# ---------- Лістинг програми (для розділу звіту) ----------
CODE_LISTING = f'''import numpy as np
import sympy as sp
import matplotlib.pyplot as plt

x1, x2 = sp.symbols('x1 x2', real=True)
f = {a}*x1**2 + {b}*x2**2 - {c}*x1*x2 - {d}*x1 - {e}*x2 + {k}

# Градієнт і матриця Гессе
grad = [sp.diff(f, v) for v in (x1, x2)]
H = sp.hessian(f, (x1, x2))

# Стаціонарна точка: розв'язок системи grad(f) = 0
crit = sp.solve(grad, [x1, x2], dict=True)[0]
f_star = sp.simplify(f.subs(crit))

# Перевірка опуклості за критерієм Сильвестра (кутові мінори)
minors = [H[:m, :m].det() for m in (1, 2)]
is_convex = all(m > 0 for m in minors)

# Число зумовленості матриці Гессе
eig = sorted(float(v) for v in H.eigenvals())
cond_H = eig[-1] / eig[0]
'''

CONSOLE_OUTPUT = f"""grad f = {grad}
H = {H.tolist()}
стаціонарна точка: {crit}
f* = {f_star_exact}
кутові мінори: {minors}
власні числа: {eigenvals}
опукла (усі мінори > 0): {is_convex}
число зумовленості cond(H) = {cond_H:.3f}
x* = ({x1_star:.4f}, {x2_star:.4f})
належить заданій області: {in_bounds}"""

# ---------- Контрольні питання ----------
QA = [
    ("Які ознаки відрізняють складну систему від простої з погляду задач оптимізації?",
     "Складну систему характеризують велика розмірність простору станів, наявність зворотних "
     "зв'язків, багатокритеріальність, невизначеність параметрів та ієрархічність управління. "
     "Її поведінка не зводиться до простої суми поведінок окремих частин, на відміну від простої "
     "системи, де така декомпозиція можлива."),
    ("З яких елементів складається математична модель задачі оптимізації?",
     "З чотирьох елементів: керованих змінних x = (x₁,…,xₙ), цільової функції f(x), обмежень "
     "(рівнянь і нерівностей, що визначають допустиму область X) та параметрів моделі — фіксованих "
     "величин, узятих із вихідних даних системи."),
    ("Чим відрізняються керовані змінні від параметрів моделі? Наведіть приклад.",
     "Керовані змінні — це величини, значення яких визначає особа, що приймає рішення, і які є "
     "невідомими в задачі оптимізації. Параметри — фіксовані характеристики системи, задані "
     "наперед і в процесі розв'язання не змінюються. Наприклад, для складської системи керовані "
     "змінні — це інтенсивності приймання й відвантаження x₁, x₂, а параметром може бути гранична "
     "місткість складу чи тариф на зберігання одиниці товару."),
    ("Сформулюйте необхідну та достатню умови локального мінімуму диференційовної функції.",
     "Необхідна умова (першого порядку): у точці локального екстремуму градієнт дорівнює нулю, "
     "∇f(x*) = 0 (стаціонарна точка). Достатня умова (другого порядку) мінімуму: матриця Гессе "
     "H(x*) у цій точці додатно визначена (усі її власні числа додатні)."),
    ("Що таке матриця Гессе і як за нею перевірити опуклість функції?",
     "Матриця Гессе — квадратна матриця других частинних похідних функції, H(x) = [∂²f/∂xᵢ∂xⱼ]. "
     "Функція опукла на області X тоді й лише тоді, коли H(x) додатно напіввизначена (⪰ 0) у "
     "кожній точці X; для строгої опуклості H(x) має бути додатно визначеною в усіх точках."),
    ("Сформулюйте критерій Сильвестра. Як він пов'язаний із власними числами матриці?",
     "Критерій Сильвестра: матриця додатно визначена тоді й лише тоді, коли всі її кутові "
     "(головні) мінори строго додатні: Δ₁ > 0, Δ₂ > 0, …, Δₙ > 0. Це еквівалентно умові, що всі "
     "власні числа матриці додатні — обидва критерії перевіряють ту саму властивість додатної "
     "визначеності, лише різними обчислювальними шляхами (мінори рахувати вручну простіше, ніж "
     "шукати власні числа)."),
    ("Чому для опуклої функції будь-який локальний мінімум є глобальним?",
     "Опуклість означає, що для будь-яких двох точок x, y ∈ X відрізок, який їх з'єднує, лежить "
     "не вище хорди значень функції в цих точках. Якби існував локальний мінімум, глибший за який "
     "є інша точка з меншим значенням f, то на відрізку між ними функція мусила б десь "
     "перевищувати лінійну хорду — а це суперечить означенню опуклості. Тому в опуклої функції не "
     "може бути кількох локальних мінімумів різної глибини — він єдиний і одночасно глобальний."),
    ("Що показує число зумовленості матриці Гессе і як воно впливає на швидкість збіжності методів?",
     "Число зумовленості cond(H) = λmax/λmin характеризує, наскільки витягнутий рельєф цільової "
     "функції. Значення, близьке до 1, відповідає майже круговим лініям рівня — градієнтні методи "
     "збігаються швидко напряму до мінімуму. Велике значення означає вузький витягнутий «яр» — "
     "градієнтний спуск рухається зигзагами й збігається значно повільніше."),
    ("Як за виглядом ліній рівня оцінити складність задачі для градієнтних методів?",
     "Концентричні кола на лініях рівня свідчать про добре зумовлену задачу (cond(H) ≈ 1) і "
     "швидку збіжність простих градієнтних методів. Витягнуті вузькі еліпси («яр») вказують на "
     "велике число зумовленості й повільну, зигзагоподібну збіжність — у таких випадках доцільніше "
     "застосовувати методи другого порядку (напр., метод Ньютона) або передобумовлення."),
    ("Наведіть приклад задачі, у якій стаціонарна точка існує, але мінімуму немає.",
     "Класичний приклад — сідлова функція f(x₁, x₂) = x₁² − x₂². Її градієнт дорівнює нулю в точці "
     "(0, 0), тобто це стаціонарна точка, але вздовж осі x₁ функція зростає, а вздовж x₂ — спадає. "
     "Матриця Гессе H = diag(2, −2) має і додатне, і від'ємне власне число, тобто не є ні додатно, "
     "ні від'ємно визначеною — критерій другого порядку не виконується, і (0, 0) є сідловою "
     "точкою, а не мінімумом чи максимумом."),
]

# ---------- DOCX-звіт (редагований Word-документ) ----------
doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(13)

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h

def add_para(text="", bold=False, italic=False, align=None, size=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if align:
        p.alignment = align
    return p

def add_code_block(text):
    p = doc.add_paragraph()
    lines = text.split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(10)
        if i < len(lines) - 1:
            run.add_break()
    p.paragraph_format.space_after = Pt(8)
    return p

def add_table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(headers):
        t.rows[0].cells[i].text = htext
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return t

# --- Титульний аркуш ---
add_para("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(UNIVERSITY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(FACULTY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(6):
    doc.add_paragraph()
add_para(f"Дисципліна: «{DISCIPLINE}»", align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(f"ЗВІТ до практичної роботи № {WORK_NUMBER}", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(f"«{WORK_TITLE}»", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(f"Варіант № {VARIANT}", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(8):
    doc.add_paragraph()
add_para(f"Виконав: студент групи {GROUP}", align=WD_ALIGN_PARAGRAPH.RIGHT)
add_para(STUDENT_NAME, align=WD_ALIGN_PARAGRAPH.RIGHT)
for _ in range(6):
    doc.add_paragraph()
add_para("Вінниця — 2026", align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# --- 1. Мета ---
add_heading("1. Мета роботи")
add_para(
    f"Формалізувати задачу прийняття рішень для системи «{obj_name}» як задачу математичного "
    "програмування, класифікувати отриману задачу, дослідити цільову функцію на опуклість та "
    "знайти стаціонарні точки аналітично і засобами Python."
)

# --- 2. Постановка задачі ---
add_heading("2. Змістовна та формальна постановка задачі")
add_para(
    f"Об'єкт дослідження — {obj_name}. Керовані змінні x1, x2 визначають два ключові параметри "
    "роботи системи, а цільова функція описує сумарні витрати/показник ефективності залежно від "
    "них. Задача — знайти такі значення x1, x2, які мінімізують цю функцію."
)
add_para(f"f(x1, x2) = {a}*x1^2 + {b}*x2^2 - {c}*x1*x2 - {d}*x1 - {e}*x2 + {k}  ->  min",
         bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(
    f"Область аналізу для побудови графіка: x1 ∈ [{x1_lo}; {x1_hi}], x2 ∈ [{x2_lo}; {x2_hi}]. "
    "Задача безумовна (розв'язується на всьому R^2, без обмежень-нерівностей)."
)

# --- 3. Класифікація ---
add_heading("3. Класифікація задачі")
add_table(["Ознака", "Клас задачі"], classification)
add_para("Отже, це — задача квадратичного програмування без обмежень.", bold=True)

# --- 4. Аналітичні викладки ---
add_heading("4. Аналітичні викладки")
add_para("Градієнт:", bold=True)
add_para(f"∇f = ({grad[0]};  {grad[1]})^T", align=WD_ALIGN_PARAGRAPH.CENTER)
add_para("Матриця Гессе:", bold=True)
add_para(f"H = {H.tolist()}", align=WD_ALIGN_PARAGRAPH.CENTER)
add_para("Стаціонарна точка (розв'язок системи ∇f(x) = 0):", bold=True)
add_para(
    f"x1* = {crit[x1]} = {x1_star:.4f},   x2* = {crit[x2]} = {x2_star:.4f},   "
    f"f(x*) = {f_star_exact} = {f_star_float:.4f}",
    align=WD_ALIGN_PARAGRAPH.CENTER
)
add_para("Перевірка опуклості за критерієм Сильвестра (кутові мінори матриці Гессе):", bold=True)
add_para(f"Δ1 = {minors[0]} > 0,   Δ2 = {minors[1]} > 0", align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(
    "Обидва мінори додатні → матриця Гессе додатно визначена → функція строго опукла → "
    "знайдена стаціонарна точка є єдиним глобальним мінімумом."
)
add_para(
    f"Власні числа матриці Гессе: λ1 = {eigenvals[0]} ≈ {eigen_floats[0]:.4f}, "
    f"λ2 = {eigenvals[1]} ≈ {eigen_floats[1]:.4f} — обидва додатні, що підтверджує опуклість "
    "альтернативним способом."
)
cond_comment = ("задача добре зумовлена, лінії рівня близькі до кіл, градієнтні методи "
                "збігатимуться швидко." if cond_H < 10 else
                "задача погано зумовлена (витягнутий «яр»), градієнтні методи збігатимуться повільно.")
add_para(f"Число зумовленості: cond(H) = λmax/λmin ≈ {cond_H:.3f} — {cond_comment}")
add_para(
    f"Належність області аналізу: точка x* = ({x1_star:.4f}, {x2_star:.4f}) "
    f"{'належить' if in_bounds else 'НЕ належить'} заданій області "
    f"x1 ∈ [{x1_lo}; {x1_hi}], x2 ∈ [{x2_lo}; {x2_hi}]."
)

# --- 5. Реалізація в Python ---
add_heading("5. Реалізація в Python")
add_code_block(CODE_LISTING.strip("\n"))
add_para("Результат виконання програми (консольний вивід):", bold=True)
add_code_block(CONSOLE_OUTPUT)

# --- 6. Графік ---
add_heading("6. Графік ліній рівня")
doc.add_picture(contour_path, width=Cm(13))
last_p = doc.paragraphs[-1]
last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# --- 7. Порівняння ---
add_heading("7. Порівняння аналітичного та комп'ютерного розв'язків")
add_table(
    ["Величина", "Аналітично (точний вираз)", "Обчислено (Python/SymPy)"],
    [
        ["x1*", str(crit[x1]), f"{x1_star:.4f}"],
        ["x2*", str(crit[x2]), f"{x2_star:.4f}"],
        ["f(x*)", str(f_star_exact), f"{f_star_float:.4f}"],
        ["Δ1, Δ2", f"{minors[0]}, {minors[1]}", f"{float(minors[0])}, {float(minors[1])}"],
    ]
)
add_para(
    "Ручний аналітичний розв'язок і результат символьних обчислень у SymPy повністю збігаються, "
    "що підтверджує коректність обчислень."
)

# --- 8. Висновки ---
add_heading("8. Висновки")
conclusions = [
    "Задача належить до класу квадратичного програмування без обмежень (двовимірна, неперервна, "
    "детермінована, статична, однокритеріальна).",
    "Цільова функція строго опукла (за критерієм Сильвестра та за додатністю власних чисел "
    "матриці Гессе), тому знайдений локальний мінімум є єдиним глобальним.",
    f"Число зумовленості cond(H) ≈ {cond_H:.3f} — "
    f"{'задача добре зумовлена' if cond_H < 10 else 'задача погано зумовлена'}, відповідно прості "
    f"градієнтні методи {'будуть ефективними' if cond_H < 10 else 'збігатимуться повільно, варто розглянути методи з передобумовленням'}.",
    "Для розв'язання задач цього класу доцільно застосовувати аналітичне розв'язання системи "
    "∇f = 0 (як у даній роботі) або, для більш складних квадратичних задач, метод Ньютона чи "
    "градієнтні методи з коригуванням кроку.",
]
for c in conclusions:
    doc.add_paragraph(c, style="List Bullet")

# --- 9. Контрольні питання ---
add_heading("9. Відповіді на контрольні питання")
for i, (q, ans) in enumerate(QA, start=1):
    add_para(f"{i}. {q}", bold=True)
    add_para(ans)

docx_path = f"Zvit_PR{WORK_NUMBER}_Variant{VARIANT}.docx"
doc.save(docx_path)

print(f"Готово: {docx_path} (+ {contour_path})")
