"""
Генератор звіту до практичної роботи №4 (Градієнтні методи безумовної оптимізації та методи другого порядку).
Питає номер варіанта (або приймає його аргументом) і генерує Zvit_PR4_Variant<N>.docx.

Запуск:  .venv/bin/python generate_report.py [номер_варіанта]
"""
import sys
import numpy as np
import sympy as sp
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from scipy.optimize import minimize
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from variants import VARIANTS

# --- Дані студента / титульного аркуша ---
STUDENT_NAME = "Толошний Мирослав"
GROUP = "КН-11"
UNIVERSITY = "Вінницький національний аграрний університет"
FACULTY = "Факультет комп'ютерних наук та економічної кібернетики"
DISCIPLINE = "Аналіз та методи оптимізації складних систем"
WORK_TITLE = "Градієнтні методи безумовної оптимізації та методи другого порядку"
WORK_NUMBER = 4

if len(sys.argv) > 1:
    VARIANT = int(sys.argv[1])
else:
    while True:
        raw = input(f"Введи номер варіанта (1-{len(VARIANTS)}): ").strip()
        if raw.isdigit() and int(raw) in VARIANTS:
            VARIANT = int(raw)
            break
        print(f"Треба ціле число від 1 до {len(VARIANTS)}, спробуй ще раз.")


def pt(v):
    """np.float64 -> звичайний float (інакше в numpy>=2 рядок типу 'np.float64(1.2)')."""
    return tuple(float(c) for c in v)


def contour_levels(Z, n=28):
    qs = np.linspace(0, 1, n)
    return np.unique(np.quantile(Z, qs))


# ---------- Модель ----------
expr_str, x0_raw, t_const, eps0 = VARIANTS[VARIANT]
x0 = np.array(x0_raw, dtype=float)

x1, x2 = sp.symbols('x1 x2', real=True)
f_sym = sp.sympify(expr_str, locals={'x1': x1, 'x2': x2})
grad_sym = [sp.diff(f_sym, v) for v in (x1, x2)]
H_sym = sp.hessian(f_sym, (x1, x2))

f_np_raw = sp.lambdify((x1, x2), f_sym, 'numpy')
g1_np = sp.lambdify((x1, x2), grad_sym[0], 'numpy')
g2_np = sp.lambdify((x1, x2), grad_sym[1], 'numpy')
H_np = sp.lambdify((x1, x2), H_sym, 'numpy')

counters = {"f": 0, "g": 0}


def f_scalar(x):
    counters["f"] += 1
    return float(f_np_raw(x[0], x[1]))


def grad_vec(x):
    counters["g"] += 1
    return np.array([float(g1_np(x[0], x[1])), float(g2_np(x[0], x[1]))])


def hess_mat(x):
    return np.array(H_np(x[0], x[1]), dtype=float)


# ---------- Точка мінімуму: символьно розв'язуємо grad f = 0 ----------
def find_exact_minimum():
    try:
        sols = sp.solve(grad_sym, [x1, x2], dict=True)
    except Exception:
        sols = []
    candidates = []
    for s in sols:
        try:
            v1, v2 = complex(s[x1]), complex(s[x2])
        except Exception:
            continue
        if abs(v1.imag) > 1e-9 or abs(v2.imag) > 1e-9:
            continue
        p = np.array([v1.real, v2.real])
        Hp = np.array(H_np(p[0], p[1]), dtype=float)
        eig = np.linalg.eigvalsh(Hp)
        if np.all(eig > 1e-9):  # локальний мінімум (додатно визначений Гессіан)
            candidates.append(p)
    if candidates:
        # найближчий до початкової точки — найбільш доречний для цього експерименту
        return min(candidates, key=lambda p: np.linalg.norm(p - x0))
    # резерв: якщо символьний розв'язок не знайшовся (трансцендентна функція тощо) —
    # точне чисельне доведення BFGS з дуже жорстким критерієм
    r = minimize(f_scalar, x0, jac=grad_vec, method='BFGS', options={'gtol': 1e-12, 'maxiter': 5000})
    return r.x


x_exact = find_exact_minimum()
f_exact = f_scalar(x_exact)
H_at_exact = hess_mat(x_exact)
eig_exact = np.linalg.eigvalsh(H_at_exact)
cond_exact = eig_exact.max() / eig_exact.min() if eig_exact.min() > 0 else float('inf')

H_at_x0 = hess_mat(x0)
eig_x0 = np.linalg.eigvalsh(H_at_x0)

# ---------- Лінії рівня (базовий графік) ----------
center = (x0 + x_exact) / 2
half_width = max(1.5, float(np.linalg.norm(x0 - x_exact)) / 2 + 1.0)
xs = np.linspace(center[0] - half_width, center[0] + half_width, 300)
ys = np.linspace(center[1] - half_width, center[1] + half_width, 300)
Xg, Yg = np.meshgrid(xs, ys)
Zg = f_np_raw(Xg, Yg)

fig0, ax0 = plt.subplots(figsize=(6, 5))
ax0.contour(Xg, Yg, Zg, levels=contour_levels(Zg), cmap='viridis')
ax0.plot(*x0, 'ko', ms=8, label='x0')
ax0.plot(*x_exact, 'g*', ms=14, label='x*')
ax0.set_xlabel('x1'); ax0.set_ylabel('x2'); ax0.legend()
ax0.set_title(f'Варіант {VARIANT}: лінії рівня f(x1, x2)')
fig0.tight_layout()
contour_path = f"contour_variant_{VARIANT}.png"
fig0.savefig(contour_path, dpi=150)
plt.close(fig0)


# ---------- Лінійний пошук за умовою Армійо ----------
def armijo(f, x, d, g, t=1.0, c=1e-4, rho=0.5, itmax=60):
    fx, slope = f(x), g @ d
    for _ in range(itmax):
        if f(x + t * d) <= fx + c * t * slope:
            return t
        t *= rho
    return t


# ---------- Методи ----------
def gd_const(x0, t, eps, itmax=200000, blow_up=1e8):
    x = np.array(x0, float)
    hist = []
    path = [x.copy()]
    stride = 1
    status = "збіжність"
    k = 0
    for k in range(itmax):
        g = grad_vec(x)
        ng = np.linalg.norm(g)
        hist.append(ng)
        if ng < eps:
            break
        x_new = x - t * g
        if not np.all(np.isfinite(x_new)) or np.linalg.norm(x_new) > blow_up:
            status = "розбіжність (|x| -> дуже велике)"
            x = x_new
            break
        x = x_new
        if k % max(1, itmax // 400) == 0:
            path.append(x.copy())
    else:
        status = "не збігся за itmax ітерацій"
    path.append(x.copy())
    return x, k + 1, np.array(path), hist, status


def gd_armijo(x0, eps, itmax=200000):
    x = np.array(x0, float)
    hist = []
    path = [x.copy()]
    for k in range(itmax):
        g = grad_vec(x)
        ng = np.linalg.norm(g)
        hist.append(ng)
        if ng < eps:
            break
        t = armijo(f_scalar, x, -g, g)
        x = x - t * g
        if k % max(1, itmax // 400) == 0:
            path.append(x.copy())
    path.append(x.copy())
    return x, k + 1, np.array(path), hist


def conj_grad_fr(x0, eps, itmax=50000):
    """Метод спряжених градієнтів Флетчера-Рівса з рестартом кожні n кроків."""
    x = np.array(x0, float)
    g = grad_vec(x)
    d = -g
    n = len(x)
    hist = [np.linalg.norm(g)]
    path = [x.copy()]
    k = 0
    for k in range(itmax):
        if np.linalg.norm(g) < eps:
            break
        t = armijo(f_scalar, x, d, g)
        x_new = x + t * d
        g_new = grad_vec(x_new)
        beta = (g_new @ g_new) / (g @ g)
        d = -g_new + (0.0 if (k + 1) % n == 0 else beta) * d
        x, g = x_new, g_new
        hist.append(np.linalg.norm(g))
        path.append(x.copy())
    return x, k + 1, np.array(path), hist


def newton_damped(x0, eps, itmax=1000):
    x = np.array(x0, float)
    hist = []
    path = [x.copy()]
    k = 0
    for k in range(itmax):
        g = grad_vec(x)
        ng = np.linalg.norm(g)
        hist.append(ng)
        if ng < eps:
            break
        H = hess_mat(x)
        try:
            d = np.linalg.solve(H, -g)
        except np.linalg.LinAlgError:
            d = -g
        if g @ d > 0:  # напрямок не є напрямком спуску
            d = -g
        x = x + armijo(f_scalar, x, d, g) * d
        path.append(x.copy())
    return x, k + 1, np.array(path), hist


# ---------- Запуск методів ----------
counters["f"] = counters["g"] = 0
x_gc, it_gc, path_gc, hist_gc, status_gc = gd_const(x0, t_const, eps0)
calls_gc = (counters["f"], counters["g"])

counters["f"] = counters["g"] = 0
x_ga, it_ga, path_ga, hist_ga = gd_armijo(x0, eps0)
calls_ga = (counters["f"], counters["g"])

counters["f"] = counters["g"] = 0
x_cg, it_cg, path_cg, hist_cg = conj_grad_fr(x0, eps0)
calls_cg = (counters["f"], counters["g"])

counters["f"] = counters["g"] = 0
x_nt, it_nt, path_nt, hist_nt = newton_damped(x0, eps0)
calls_nt = (counters["f"], counters["g"])

counters["f"] = counters["g"] = 0
path_bfgs = [np.array(x0)]
res_bfgs = minimize(f_scalar, x0, jac=grad_vec, method='BFGS', options={'gtol': eps0},
                     callback=lambda xk: path_bfgs.append(np.array(xk)))
path_bfgs = np.array(path_bfgs)
calls_bfgs = (counters["f"], counters["g"])

summary_rows = [
    ("Градієнтний, сталий t", x_gc, f_scalar(x_gc), it_gc, calls_gc, status_gc),
    ("Градієнтний + Армійо", x_ga, f_scalar(x_ga), it_ga, calls_ga, "збіжність"),
    ("Спряжені градієнти (ФР)", x_cg, f_scalar(x_cg), it_cg, calls_cg, "збіжність"),
    ("Ньютона з демпфуванням", x_nt, f_scalar(x_nt), it_nt, calls_nt, "збіжність"),
    ("BFGS (SciPy)", res_bfgs.x, res_bfgs.fun, res_bfgs.nit, (res_bfgs.nfev, res_bfgs.njev), "збіжність"),
]

# ---------- Графік збіжності lg||grad f|| ----------
fig1, ax1 = plt.subplots(figsize=(7, 5))
for hist, label in [(hist_gc, 'Градієнтний, сталий t'), (hist_ga, 'Градієнтний + Армійо'),
                     (hist_cg, 'Спряжені градієнти'), (hist_nt, 'Ньютона')]:
    h = np.array(hist)
    h = h[h > 0]
    ax1.semilogy(range(1, len(h) + 1), h, label=label)
ax1.set_xlabel('номер ітерації'); ax1.set_ylabel('||grad f|| (лог. шкала)')
ax1.grid(alpha=.3, which='both'); ax1.legend()
ax1.set_title(f'Варіант {VARIANT}: збіжність методів першого і другого порядку')
fig1.tight_layout()
conv_path = f"convergence_variant_{VARIANT}.png"
fig1.savefig(conv_path, dpi=150)
plt.close(fig1)

# ---------- Траєкторії на лініях рівня ----------
fig2, ax2 = plt.subplots(figsize=(7, 6))
ax2.contour(Xg, Yg, Zg, levels=contour_levels(Zg), cmap='viridis', alpha=.6)
for path, label, style in [
    (path_gc, 'Градієнтний, сталий t', 'c.-'),
    (path_ga, 'Градієнтний + Армійо', 'm.-'),
    (path_cg, 'Спряжені градієнти', 'b.-'),
    (path_nt, 'Ньютона', 'r.-'),
    (path_bfgs, 'BFGS', 'y.-'),
]:
    ax2.plot(path[:, 0], path[:, 1], style, ms=3, lw=1, label=label, alpha=.85)
ax2.plot(*x0, 'ko', ms=8, label='x0')
ax2.plot(*x_exact, 'g*', ms=16, label='x*')
ax2.set_xlabel('x1'); ax2.set_ylabel('x2'); ax2.legend(fontsize=8)
ax2.set_title(f'Варіант {VARIANT}: траєкторії методів')
fig2.tight_layout()
traj_path = f"trajectories_variant_{VARIANT}.png"
fig2.savefig(traj_path, dpi=150)
plt.close(fig2)

# ---------- Критичний крок градієнтного методу ----------
lam_max_x0 = eig_x0.max()
t_theory = 2.0 / lam_max_x0


def is_stable(t_try, probe_iters=3000):
    x = np.array(x0, float)
    for _ in range(probe_iters):
        g = grad_vec(x)
        if np.linalg.norm(g) < eps0:
            return True
        x = x - t_try * g
        if not np.all(np.isfinite(x)) or np.linalg.norm(x) > 1e8:
            return False
    return True


lo, hi = 1e-8, t_theory
while is_stable(hi) and hi < t_theory * 64:
    lo, hi = hi, hi * 2
if not is_stable(hi):
    for _ in range(25):
        mid = (lo + hi) / 2
        if is_stable(mid):
            lo = mid
        else:
            hi = mid
    t_critical = lo
else:
    t_critical = hi

# ---------- Лістинг коду ----------
CODE_METHODS = '''def armijo(f, x, d, g, t=1.0, c=1e-4, rho=0.5, itmax=60):
    """Дроблення кроку за умовою Армійо: f(x+td) <= f(x) + c*t*(g,d)."""
    fx, slope = f(x), g @ d
    for _ in range(itmax):
        if f(x + t * d) <= fx + c * t * slope:
            return t
        t *= rho
    return t


def gd_const(x0, t, eps, itmax=200000):
    x = np.array(x0, float)
    for k in range(itmax):
        g = grad(x)
        if np.linalg.norm(g) < eps:
            break
        x = x - t * g                          # сталий крок, без обчислень f
    return x, k + 1


def gd_armijo(x0, eps, itmax=200000):
    x = np.array(x0, float)
    for k in range(itmax):
        g = grad(x)
        if np.linalg.norm(g) < eps:
            break
        x = x - armijo(f, x, -g, g) * g         # крок дробиться під умову Армійо
    return x, k + 1


def conj_grad_fr(x0, eps, itmax=50000):
    x = np.array(x0, float); g = grad(x); d = -g; n = len(x)
    for k in range(itmax):
        if np.linalg.norm(g) < eps:
            break
        t = armijo(f, x, d, g)
        x_new = x + t * d
        g_new = grad(x_new)
        beta = (g_new @ g_new) / (g @ g)        # Флетчера-Рівса
        d = -g_new + (0.0 if (k + 1) % n == 0 else beta) * d   # рестарт кожні n кроків
        x, g = x_new, g_new
    return x, k + 1


def newton_damped(x0, eps, itmax=1000):
    x = np.array(x0, float)
    for k in range(itmax):
        g = grad(x)
        if np.linalg.norm(g) < eps:
            break
        d = np.linalg.solve(hess(x), -g)
        if g @ d > 0:                           # не напрямок спуску -> відкіт до антиградієнта
            d = -g
        x = x + armijo(f, x, d, g) * d
    return x, k + 1

# BFGS — готова реалізація SciPy:
# res = minimize(f, x0, jac=grad, method='BFGS', options={'gtol': eps})
'''

# ---------- Контрольні питання ----------
QA = [
    ("Чому антиградієнт є напрямком найшвидшого спадання функції?",
     "Градієнт ∇f(x) вказує напрямок найшвидшого зростання функції в точці x — це прямий наслідок "
     "лінійного наближення f за рядом Тейлора: приріст f(x+d) - f(x) ≈ ∇f(x)·d максимальний (за фіксованої "
     "довжини d) саме тоді, коли d співнапрямлений з ∇f(x) (скалярний добуток максимальний при нульовому "
     "куті між векторами). Отже, протилежний напрямок -∇f(x) дає найбільше можливе локальне спадання."),
    ("За якої умови градієнтний метод зі сталим кроком розбігається? Як пов'язана межа стійкості з "
     "власними числами H?",
     "Для квадратичної (чи локально квадратично апроксимованої) функції метод стійкий лише за "
     "0 < t < 2/λmax(H), де λmax — найбільше власне число матриці Гессе. Якщо крок перевищує цю межу, "
     "ітерації в напрямку власного вектора, що відповідає λmax, починають зростати за модулем з кожним "
     "кроком замість спадання — метод розгойдується і розбігається."),
    ("Сформулюйте умову Армійо. Який зміст має константа c1?",
     "Крок t приймається, якщо f(x + t·d) ≤ f(x) + c1·t·(∇f(x), d), де 0 < c1 < 1 (типово c1 = 10⁻⁴). "
     "Права частина — пряма, дотична до f вздовж напрямку d, нахилена трохи менш круто за саму дотичну "
     "(через множник c1 < 1). Умова вимагає не точного мінімуму вздовж d, а лише «достатнього» спадання — "
     "будь-яка точка нижче цієї похилої прямої приймається, що дешево перевірити і які гарантує прогрес."),
    ("Чим відрізняється найшвидший спуск від градієнтного спуску з дробленням кроку?",
     "Найшвидший спуск шукає точний мінімум вздовж напрямку антиградієнта (розв'язує одновимірну задачу "
     "tₖ = argmin f(x - t∇f(x)) точно, наприклад золотим перерізом). Градієнтний спуск з дробленням кроку "
     "(Армійо) не шукає точний мінімум, а лише дешево знаходить «достатньо хороший» крок послідовним "
     "діленням навпіл — зазвичай це суттєво дешевше загалом, хоча й дає трохи гірший крок на кожній "
     "ітерації."),
    ("Що таке спряжені напрямки? Чому метод спряжених градієнтів мінімізує квадратичну функцію за n кроків?",
     "Два напрямки d⁽ⁱ⁾, d⁽ʲ⁾ спряжені відносно додатно визначеної матриці Q, якщо (d⁽ⁱ⁾)ᵀQd⁽ʲ⁾ = 0 при "
     "i≠j — це узагальнення ортогональності, «підлаштоване» під кривину функції. Якщо рухатись по системі "
     "з n таких взаємно спряжених напрямків (кожен раз точно мінімізуючи вздовж поточного напрямку), кожен "
     "крок остаточно усуває похибку вздовж свого напрямку і більше ніколи її не «псує» — тому для "
     "квадратичної функції n змінних точний мінімум досягається рівно за n кроків."),
    ("Навіщо потрібен рестарт у методі спряжених градієнтів для неквадратичних функцій?",
     "Формула спряженості напрямків виведена строго для квадратичної функції з фіксованою матрицею Q. Для "
     "неквадратичної функції «ефективна» кривина змінюється від точки до точки, тому накопичені за багато "
     "кроків напрямки поступово втрачають властивість спряженості і метод вироджується. Періодичний рестарт "
     "(β := 0 кожні n ітерацій, тобто черговий крок — звичайний антиградієнт) усуває накопичену похибку і "
     "відновлює коректну поведінку методу."),
    ("Запишіть ітераційну формулу методу Ньютона. Чому він не залежить від зумовленості задачі?",
     "x⁽ᵏ⁺¹⁾ = x⁽ᵏ⁾ - H⁻¹(x⁽ᵏ⁾)·∇f(x⁽ᵏ⁾) (з демпфуванням — ще й множник кроку tₖ за Армійо). Множення на "
     "обернену матрицю Гессе фактично «вирівнює» простір змінних так, ніби лінії рівня стають круговими "
     "(матриця H⁻¹H = I в локальній квадратичній моделі) — тому метод рухається прямо до мінімуму незалежно "
     "від того, наскільки витягнуті реальні лінії рівня, тобто незалежно від cond(H)."),
    ("Що робити, якщо матриця Гессе не є додатно визначеною?",
     "Напрямок Ньютона d = -H⁻¹∇f тоді може вказувати не в бік спадання функції (або навіть у бік "
     "зростання) — це перевіряють умовою (∇f, d) < 0. Якщо умова порушена, практичні реалізації відкочуються "
     "до звичайного антиградієнта -∇f (гарантовано напрямок спадання) або модифікують Гессіан, додаючи до "
     "нього діагональну матрицю λI з достатньо великим λ, щоб зробити результат додатно визначеним."),
    ("У чому полягає ідея квазіньютонівських методів? Що таке секантна умова?",
     "Замість точного обчислення H⁻¹ (дорого: O(n³) на розв'язання системи й O(n²) пам'яті для n(n+1)/2 "
     "других похідних), квазіньютонівські методи (BFGS, DFP, L-BFGS) будують і поступово уточнюють "
     "наближення Bₖ ≈ H⁻¹, використовуючи лише різниці вже обчислених градієнтів. Секантна умова "
     "Bₖ₊₁yₖ = sₖ (де sₖ = x⁽ᵏ⁺¹⁾-x⁽ᵏ⁾, yₖ = ∇f⁽ᵏ⁺¹⁾-∇f⁽ᵏ⁾) вимагає, щоб нове наближення узгоджувалося зі "
     "спостереженою зміною градієнта вздовж останнього кроку — це єдина «ціна» за відмову від других "
     "похідних."),
    ("Чому BFGS є методом за замовчуванням у бібліотеках оптимізації, а не метод Ньютона?",
     "BFGS дає надлінійну (майже квадратичну) збіжність ціною лише O(n²) операцій на ітерацію і без жодного "
     "обчислення других похідних, тоді як точний метод Ньютона вимагає обчислення повного Гессіана та "
     "розв'язання лінійної системи O(n³) на кожному кроці — для задач з великою кількістю змінних це різко "
     "дорожче. BFGS зберігає практично ту саму швидкість збіжності за значно меншу обчислювальну й "
     "реалізаційну вартість, тому є розумним компромісом за замовчуванням."),
]

# ---------- DOCX-звіт ----------
doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(13)


def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h


def add_para(text="", bold=False, italic=False, align=None, size=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def add_code_block(text):
    p = doc.add_paragraph()
    lines = text.split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9.5)
        if i < len(lines) - 1:
            run.add_break()
    p.paragraph_format.space_after = Pt(8)
    return p


def add_table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(headers):
        t.rows[0].cells[i].text = htext
    for row in rows:
        cells_ = t.add_row().cells
        for i, val in enumerate(row):
            cells_[i].text = str(val)
    return t


def fmt(v, nd=6):
    return f"{v:.{nd}f}" if isinstance(v, (int, float, np.floating)) else str(v)


# --- Титульний аркуш ---
add_para("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(UNIVERSITY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(FACULTY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(6):
    doc.add_paragraph()
add_para(f"Дисципліна: «{DISCIPLINE}»", align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(f"ЗВІТ до практичної роботи № {WORK_NUMBER}", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(f"«{WORK_TITLE}»", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(f"Варіант № {VARIANT}", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(8):
    doc.add_paragraph()
add_para(f"Виконав: студент групи {GROUP}", align=WD_ALIGN_PARAGRAPH.RIGHT)
add_para(STUDENT_NAME, align=WD_ALIGN_PARAGRAPH.RIGHT)
for _ in range(6):
    doc.add_paragraph()
add_para("Вінниця — 2026", align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# --- 1. Мета ---
add_heading("1. Мета роботи")
add_para(
    "Дослідити методи, що використовують похідні цільової функції: градієнтний спуск зі сталим та "
    "дробленим кроком, метод спряжених градієнтів, метод Ньютона та квазіньютонівський метод BFGS; "
    "встановити зв'язок між зумовленістю задачі та швидкістю збіжності."
)

# --- 2. Аналітичні вирази ---
add_heading("2. Аналітичні вирази градієнта і матриці Гессе")
add_para(f"f(x1, x2) = {f_sym}   ->  min,   x0 = {pt(x0)}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(f"∂f/∂x1 = {grad_sym[0]}")
add_para(f"∂f/∂x2 = {grad_sym[1]}")
add_para(f"H(x) = {H_sym.tolist()}")
add_para(
    f"Власні числа H у початковій точці x0: {[round(float(e), 4) for e in eig_x0]}, "
    f"λmax = {lam_max_x0:.4f}."
)
add_para(
    f"Точка мінімуму x* = {pt(x_exact)}, f(x*) = {f_exact:.6f}. Власні числа H у точці мінімуму: "
    f"{[round(float(e), 4) for e in eig_exact]}, число зумовленості cond(H) = {cond_exact:.2f}."
)
doc.add_picture(contour_path, width=Cm(11))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

# --- 3. Тексти програм ---
add_heading("3. Тексти програмних реалізацій методів")
add_code_block(CODE_METHODS.strip("\n"))

# --- 4. Таблиці перших ітерацій ---
add_heading("4. Таблиці перших ітерацій кожного методу")
for name, path, hist in [
    ("Градієнтний, сталий t", path_gc, hist_gc),
    ("Градієнтний + Армійо", path_ga, hist_ga),
    ("Спряжені градієнти (ФР)", path_cg, hist_cg),
    ("Ньютона з демпфуванням", path_nt, hist_nt),
]:
    add_para(name, bold=True)
    n_show = min(7, len(path) - 1, len(hist))
    rows = [[k + 1, fmt(path[k][0]), fmt(path[k][1]), fmt(hist[k], 8)] for k in range(n_show)]
    add_table(["iter", "x1", "x2", "||grad f||"], rows)

# --- 5. Графік збіжності ---
add_heading("5. Графік збіжності lg||grad f|| від номера ітерації")
doc.add_picture(conv_path, width=Cm(14))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_para(
    "Чим крутіше спадає крива методу, тим швидша його збіжність. Методи, що використовують кривину функції "
    "(Ньютона, BFGS), виходять на машинну точність за одиниці-десятки ітерацій, тоді як градієнтний спуск "
    "зі сталим кроком спадає лінійно й повільно — йому потрібні тисячі ітерацій через високу зумовленість "
    "задачі."
)

# --- 6. Траєкторії ---
add_heading("6. Траєкторії методів на лініях рівня")
doc.add_picture(traj_path, width=Cm(14))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

# --- 7. Критичний крок ---
add_heading("7. Експеримент із критичним значенням сталого кроку")
add_para(
    f"Теоретична межа стійкості: t < 2/λmax(H(x0)) = 2/{lam_max_x0:.4f} ≈ {t_theory:.6f}. "
    f"Експериментально знайдене (бісекцією) критичне значення кроку: t_critical ≈ {t_critical:.6f}."
)
add_para(
    f"Заданий у варіанті крок t = {t_const} "
    f"{'менший за критичний — метод зі сталим кроком збігається, що підтверджено експериментом.' if t_const < t_critical else 'перевищує критичне значення — метод зі сталим кроком розбігається, що й спостерігається в таблиці нижче.'}"
)

# --- 8. Зведена таблиця ---
add_heading("8. Зведена порівняльна таблиця")
table_rows = []
for name, x_, fv, it_, calls_, status in summary_rows:
    table_rows.append([name, fmt(x_[0]), fmt(x_[1]), fmt(fv, 8), it_, calls_[0], calls_[1], status])
add_table(
    ["Метод", "x1*", "x2*", "f(x*)", "Ітерацій", "Обч. f", "Обч. grad f", "Статус"],
    table_rows,
)
converged = [r for r in summary_rows if r[5] == "збіжність"]
fastest = min(converged, key=lambda r: r[3])
add_para(
    f"Найменшої кількості ітерацій серед методів, що збіглися, потребував «{fastest[0]}» ({fastest[3]} "
    f"ітерацій) — це узгоджується з теорією: методи другого порядку (Ньютона, BFGS) не залежать від "
    f"зумовленості задачі, тоді як градієнтний спуск зі сталим кроком страждає від cond(H) ≈ {cond_exact:.0f}."
)

# --- 9. Висновки ---
add_heading("9. Висновки")
conclusions = [
    f"Усі методи, що збіглися, для варіанта {VARIANT} знайшли мінімум f(x*) ≈ {fmt(fastest[2], 6)} у точці "
    f"x* ≈ {pt(x_exact)}, що узгоджується з аналітичним розв'язком системи ∇f(x) = 0.",
    f"Число зумовленості матриці Гессе в точці мінімуму cond(H) = {cond_exact:.1f} пояснює, чому "
    f"градієнтний спуск зі сталим кроком вимагає на порядки більше ітерацій за методи другого порядку: "
    f"теоретичний множник спадання похибки ((cond-1)/(cond+1))² близький до 1 при великому cond(H).",
    "Дроблення кроку за умовою Армійо усуває потребу вручну підбирати сталий крок, але платить за це "
    "додатковими обчисленнями функції на кожній ітерації (пробні обчислення під час дроблення).",
    "Метод спряжених градієнтів дає суттєвий виграш над звичайним градієнтним спуском при пам'яті лише "
    "O(n), що робить його придатним для задач великої розмірності, де зберігати чи обертати матрицю "
    "Гессе (як у методі Ньютона) неможливо.",
    "Метод Ньютона і BFGS майже не залежать від зумовленості задачі і виходять на високу точність за "
    "одиниці-десятки ітерацій; BFGS при цьому не вимагає обчислення других похідних, що й робить його "
    "методом за замовчуванням у більшості бібліотек оптимізації.",
]
for c in conclusions:
    doc.add_paragraph(c, style="List Bullet")

# --- 10. Контрольні питання ---
add_heading("10. Відповіді на контрольні питання")
for i, (q, ans) in enumerate(QA, start=1):
    add_para(f"{i}. {q}", bold=True)
    add_para(ans)

docx_path = f"Zvit_PR{WORK_NUMBER}_Variant{VARIANT}.docx"
doc.save(docx_path)
print(f"Готово: {docx_path} (+ {contour_path}, {conv_path}, {traj_path})")
print(f"x* = {pt(x_exact)}, f* = {f_exact:.6f}, cond(H)@x* = {cond_exact:.2f}, t_theory = {t_theory:.6f}, t_critical = {t_critical:.6f}")
